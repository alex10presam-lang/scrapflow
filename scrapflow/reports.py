"""Relatório independente: nunca altera arquivos de entrada."""

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from .pipeline import summarize


def safe_text(value):
    # Textos importados não devem se tornar fórmulas no arquivo exportado.
    return (
        "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
    )


def workbook_report(run):
    summary = summarize(run["dataset"])
    book = Workbook()
    cover = book.active
    cover.title = "Resumo"
    rows = [
        ["SCRAPFLOW", "Relatório de perdas"],
        ["Execução", run["id"]],
        ["Origem", run["source"]],
        ["Gerado em (UTC)", run["created_at"]],
        ["Unidades boas", summary["good"]],
        ["Unidades perdidas", summary["lost"]],
        ["Taxa de perda", summary["rate"] / 100],
        ["Custo estimado (BRL)", summary["cost_cents"] / 100],
        ["Critério", "Perdidas / (boas + perdidas). Custo = quantidade perdida × custo unitário."],
        [
            "Nota",
            (
                "Demonstração com dados sintéticos"
                if run["source"] == "demo"
                else "Lote importado pelo usuário"
            ),
        ],
    ]
    for row in rows:
        cover.append(row)
    cover["B7"].number_format = "0.00%"
    cover["B8"].number_format = '"R$" #,##0.00'
    detail = book.create_sheet("Perdas")
    columns = (
        "id",
        "data",
        "ordem",
        "produto",
        "descricao",
        "linha",
        "motivo",
        "quantidade",
        "custo_centavos",
    )
    detail.append(
        [
            "ID",
            "Data",
            "Ordem",
            "Produto",
            "Descrição",
            "Linha",
            "Motivo",
            "Quantidade",
            "Custo (BRL)",
        ]
    )
    for record in run["dataset"]["records"]:
        detail.append(
            [
                safe_text(record[key]) if key != "custo_centavos" else record[key] / 100
                for key in columns
            ]
        )
    for cell in list(detail.columns)[8][1:]:
        cell.number_format = '"R$" #,##0.00'
    reasons = book.create_sheet("Causas")
    reasons.append(["Motivo", "Custo estimado (BRL)"])
    for reason in summary["reasons"]:
        reasons.append([safe_text(reason["name"]), reason["cents"] / 100])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions if sheet != cover else "A1:B1"
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="142C3B")
        for index, column in enumerate(sheet.columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(
                65, max(16, max(len(str(c.value or "")) for c in column) + 3)
            )
        for row in list(sheet.rows)[1:]:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    stream = BytesIO()
    book.save(stream)
    book.close()
    stream.seek(0)
    return stream
