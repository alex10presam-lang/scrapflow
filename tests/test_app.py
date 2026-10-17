from copy import deepcopy
from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook
from lastro import ROOT, create_app, sample_data
from lastro.pipeline import summarize


class ApplicationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.app = create_app({"TESTING": True, "DATABASE": Path(self.temp.name) / "runs.db"})
        self.client = self.app.test_client()

    def files(self):
        return {
            name: (BytesIO((ROOT / "examples" / f"{name}.csv").read_bytes()), f"{name}.csv")
            for name in ("produtos", "producao", "perdas")
        }

    def test_preview_does_not_create_history(self):
        response = self.client.get("/api/dashboard")
        self.assertTrue(response.json["preview"])
        self.assertEqual(self.client.get("/api/runs").json["runs"], [])

    def test_demo_can_be_saved_reloaded_and_exported(self):
        response = self.client.post("/api/runs/demo")
        self.assertEqual(response.status_code, 201)
        run_id = response.json["id"]
        result = self.client.get(f"/api/dashboard?run={run_id}").json
        self.assertFalse(result["preview"])
        self.assertEqual(result["summary"]["lost"], 400)
        self.assertEqual(result["summary"]["good"], 12378)
        self.assertEqual(result["summary"]["cost_cents"], 747170)
        report = self.client.get(f"/api/runs/{run_id}/report")
        self.assertEqual(report.status_code, 200)
        book = load_workbook(BytesIO(report.data))
        self.assertEqual(book.sheetnames, ["Resumo", "Perdas", "Causas"])
        self.assertEqual(book["Perdas"].max_row, 37)
        self.assertEqual(book["Resumo"]["B8"].value, 7471.7)
        book.close()

    def test_upload_uses_real_pipeline(self):
        response = self.client.post("/api/runs", data=self.files())
        self.assertEqual(response.status_code, 201)
        run = self.app.extensions["runs"].get(response.json["id"])
        self.assertEqual(run["source"], "upload")
        self.assertEqual(summarize(run["dataset"])["lost"], 400)

    def test_invalid_batch_leaves_previous_run_unchanged(self):
        run_id = self.client.post("/api/runs/demo").json["id"]
        files = self.files()
        files["perdas"] = (
            BytesIO(b"id,ordem,data,quantidade,motivo\nP,UNKNOWN,2026-10-01,1,Test"),
            "perdas.csv",
        )
        response = self.client.post("/api/runs", data=files)
        self.assertEqual(response.status_code, 422)
        self.assertEqual(len(self.client.get("/api/runs").json["runs"]), 1)
        self.assertEqual(self.app.extensions["runs"].get()["id"], run_id)

    def test_partial_batch_rejected(self):
        response = self.client.post("/api/runs", data={"produtos": (BytesIO(b"x"), "test.csv")})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(self.client.get("/api/runs").json["runs"], [])

    def test_large_upload_rejected(self):
        self.app.config["MAX_CONTENT_LENGTH"] = 100
        self.assertEqual(self.client.post("/api/runs", data=self.files()).status_code, 413)

    def test_unknown_run_and_report(self):
        self.assertEqual(self.client.get("/api/dashboard?run=unknown").status_code, 404)
        self.assertEqual(self.client.get("/api/runs/unknown/report").status_code, 404)

    def test_cross_origin_rejected(self):
        response = self.client.post("/api/runs/demo", headers={"Origin": "https://example.com"})
        self.assertEqual(response.status_code, 403)

    def test_report_escapes_formula_like_text(self):
        dataset = deepcopy(sample_data())
        dataset["records"][0]["motivo"] = '=HYPERLINK("https://example.org")'
        run = self.app.extensions["runs"].save(dataset, "upload")
        report = self.client.get(f'/api/runs/{run["id"]}/report')
        book = load_workbook(BytesIO(report.data))
        cell = book["Perdas"]["G2"]
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.value.startswith("'="))
        book.close()

    def test_examples_download_contains_only_synthetic_inputs(self):
        response = self.client.get("/api/examples")
        with ZipFile(BytesIO(response.data)) as archive:
            self.assertEqual(
                set(archive.namelist()), {"produtos.csv", "producao.csv", "perdas.csv"}
            )

    def test_dashboard_filters_by_line(self):
        all_data = self.client.get("/api/dashboard").json["summary"]
        filtered = self.client.get("/api/dashboard?line=Usinagem").json["summary"]
        self.assertLess(filtered["lost"], all_data["lost"])
        self.assertTrue(all(row["linha"] == "Usinagem" for row in filtered["records"]))

    def test_static_assets_and_headers(self):
        for path in ("/", "/static/app.css", "/static/app.js", "/static/mark.svg"):
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")
                response.close()
